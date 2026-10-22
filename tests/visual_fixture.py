"""Optional renderer smoke fixture. Requires existing python-docx; installs nothing."""
import argparse
import importlib.util
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def create(folder):
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    title_properties = doc.styles['Title'].element.find(qn('w:pPr'))
    if title_properties is not None:
        for border in list(title_properties.findall(qn('w:pBdr'))):
            title_properties.remove(border)
    doc.add_heading('Synthetic handover verification fixture', 0)
    doc.add_paragraph('This fixture contains invented engineering information for software tests only.')
    headings = ['Overview', 'Architecture', 'Setup', 'Workflows', 'API integration', 'Testing', 'Known issues and limitations', 'Handover']
    edit_paragraph = None
    for i, heading in enumerate(headings, 1):
        doc.add_heading(f'{i}. {heading}', 1)
        if heading == 'Architecture':
            edit_paragraph = len(doc.paragraphs)
            para = doc.add_paragraph()
            para.add_run('The client uses ').bold = True
            para.add_run('local ').italic = True
            para.add_run('drafts for this demonstration.')
        else:
            doc.add_paragraph('The synthetic component has a documented boundary and a maintainer action. '
                              'Observed tests, source claims and pending acceptance work are recorded separately. '
                              'No external network or production service is used by this fixture.')
    doc.add_paragraph('Table 1. Synthetic maintenance inventory', 'Caption')
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['ID', 'Responsibility', 'Verification scope']):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i in range(32):
        cells = table.add_row().cells
        for cell, text in zip(cells, [str(i + 1), 'Synthetic component ' + str(i + 1), 'Unit assertions; runtime acceptance pending.']):
            cell.text = text
    merged = table.add_row().cells
    merged[0].merge(merged[2]).text = 'Merged summary row. All values are synthetic.'
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    footer = sec.footer.paragraphs[0]
    footer.add_run('Synthetic fixture | Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    footer._p.append(field)
    original = folder / 'original.docx'
    doc.save(original)
    spec = importlib.util.spec_from_file_location('engine', Path(__file__).parents[1] / 'scripts/docx_engineer.py')
    engine = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(engine)
    pkg = engine.Package(original)
    # Use actual OOXML paragraph index rather than assuming direct-body indexing.
    index = next(i for i, p in enumerate(pkg.paragraphs) if engine.text_of(p).startswith('The client uses'))
    plan = {'input_sha256': pkg.digest, 'replacements': [{'paragraph': index, 'old': 'local drafts', 'new': 'session drafts'}]}
    candidate = folder / 'candidate.docx'
    result = engine.edit(pkg, plan, candidate)
    (folder / 'edit-plan.json').write_text(json.dumps(plan, indent=2))
    (folder / 'patch-result.json').write_text(json.dumps(result, indent=2))
    print(json.dumps({'original': str(original), 'candidate': str(candidate), 'paragraph': index}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, required=True)
    create(parser.parse_args().output_dir)
